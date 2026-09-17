import csv
import email
from email import policy
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import List, Dict, Any

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

from app.document_processing.ocr_engine import ocr_engine

MAX_ARCHIVE_FILES = 25
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024  # 50 MB limit
ALLOWED_ARCHIVE_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".csv", ".txt", ".eml", ".png", ".jpg", ".jpeg"}
DISALLOWED_EXECUTABLE_EXTENSIONS = {".exe", ".bat", ".cmd", ".sh", ".py", ".bin", ".dll", ".so", ".vbs", ".ps1", ".jar"}

class DocumentExtractor:
    @staticmethod
    def extract(file_path: Path, file_type: str) -> List[Dict[str, Any]]:
        """
        Extracts structured page content from various file formats.
        Returns a list of dicts: [{"page": 1, "section": "...", "content": "...", "extraction_mode": "..."}]
        """
        # Safety check: empty file
        if file_path.exists() and file_path.stat().st_size == 0:
            return [{
                "page": 1,
                "section": "Error",
                "content": f"File '{file_path.name}' is empty (0 bytes). Extraction aborted.",
                "extraction_mode": "failed"
            }]

        ext = file_type.lower().replace(".", "")
        if f".{ext}" in DISALLOWED_EXECUTABLE_EXTENSIONS or ext in ["exe", "bat", "cmd", "sh", "bin", "dll", "so", "vbs", "ps1", "jar"]:
            return [{
                "page": 1,
                "section": "Security Error",
                "content": f"Executable file type '.{ext}' is rejected for industrial safety and security.",
                "extraction_mode": "failed"
            }]

        if ext == "pdf":
            return DocumentExtractor._extract_pdf(file_path)
        elif ext in ["png", "jpg", "jpeg", "tiff", "tif", "bmp"]:
            return DocumentExtractor._extract_image(file_path)
        elif ext in ["docx", "doc"]:
            return DocumentExtractor._extract_docx(file_path)
        elif ext in ["xlsx", "xls"]:
            return DocumentExtractor._extract_xlsx(file_path)
        elif ext == "csv":
            return DocumentExtractor._extract_csv(file_path)
        elif ext in ["eml", "msg"]:
            return DocumentExtractor._extract_email(file_path)
        elif ext == "zip":
            return DocumentExtractor._extract_archive(file_path)
        elif ext in ["txt", "text", "log", "md"]:
            return DocumentExtractor._extract_text(file_path)
        else:
            return [{
                "page": 1,
                "section": "Error",
                "content": f"Unsupported document format '.{ext}'. Supported formats: PDF, DOCX, XLSX, CSV, TXT, PNG, JPG/JPEG, EML, ZIP.",
                "extraction_mode": "failed"
            }]

    @staticmethod
    def _extract_image(file_path: Path) -> List[Dict[str, Any]]:
        """Extracts text from raster image pixels using Tesseract OCR."""
        try:
            ocr_res = ocr_engine.process_scanned_document(file_path, document_title=file_path.stem)
            raw_text = ocr_res.get("extracted_text", "")
            return [{
                "page": 1,
                "section": file_path.stem,
                "content": raw_text,
                "extraction_mode": "ocr_raster_extraction",
                "ocr_confidence_pct": ocr_res.get("overall_confidence_pct", 90.0)
            }]
        except Exception as e:
            return [{
                "page": 1,
                "section": "Error",
                "content": f"Image OCR extraction failed: {str(e)}",
                "extraction_mode": "failed"
            }]

    @staticmethod
    def _extract_pdf(file_path: Path) -> List[Dict[str, Any]]:
        if PdfReader is None:
            return [{
                "page": 1,
                "section": "Error",
                "content": "pypdf package is not available in the current Python environment.",
                "extraction_mode": "failed"
            }]
        pages = []
        try:
            reader = PdfReader(str(file_path))
            if reader.is_encrypted:
                return [{
                    "page": 1,
                    "section": "Error",
                    "content": "Password-protected PDF. Document extraction aborted.",
                    "extraction_mode": "failed"
                }]

            total_text_len = 0
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                total_text_len += len(text.strip())
                lines = [line.strip() for line in text.split("\n") if line.strip()]
                section = lines[0] if lines else f"Page {idx + 1}"
                pages.append({
                    "page": idx + 1,
                    "section": section[:80],
                    "content": text,
                    "extraction_mode": "text_layer_extraction"
                })

            # If PDF text layer is empty (e.g. pure scanned image in PDF), invoke raster OCR
            if total_text_len < 20 and len(reader.pages) > 0:
                ocr_res = ocr_engine.process_scanned_document(file_path, document_title=file_path.stem)
                ocr_text = ocr_res.get("extracted_text", "")
                if ocr_text.strip():
                    return [{
                        "page": 1,
                        "section": file_path.stem,
                        "content": ocr_text,
                        "extraction_mode": "ocr_raster_extraction",
                        "ocr_confidence_pct": ocr_res.get("overall_confidence_pct", 85.0)
                    }]

        except Exception as e:
            pages.append({
                "page": 1,
                "section": "Error",
                "content": f"Failed to extract PDF: {str(e)}",
                "extraction_mode": "failed"
            })
        return pages

    @staticmethod
    def _extract_docx(file_path: Path) -> List[Dict[str, Any]]:
        if DocxDocument is None:
            return [{
                "page": 1,
                "section": "Error",
                "content": "python-docx package is not available in the current Python environment."
            }]
        try:
            doc = DocxDocument(str(file_path))
            pages = []
            current_section = "General Information"
            current_text = []
            page_num = 1
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                style_name = getattr(para.style, "name", "") if para.style else ""
                if style_name.startswith("Heading"):
                    if current_text:
                        pages.append({
                            "page": page_num,
                            "section": current_section,
                            "content": "\n".join(current_text)
                        })
                        current_text = []
                        page_num += 1
                    current_section = text
                else:
                    current_text.append(text)
            
            if current_text:
                pages.append({
                    "page": page_num,
                    "section": current_section,
                    "content": "\n".join(current_text)
                })
            return pages if pages else [{"page": 1, "section": "General", "content": ""}]
        except Exception as e:
            return [{"page": 1, "section": "Error", "content": f"DOCX extraction failed: {str(e)}"}]

    @staticmethod
    def _extract_xlsx(file_path: Path) -> List[Dict[str, Any]]:
        if openpyxl is None:
            return [{
                "page": 1,
                "section": "Error",
                "content": "openpyxl package is not available in the current Python environment."
            }]
        try:
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            pages = []
            for sheet_idx, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_strs = [str(cell) for cell in row if cell is not None]
                    if row_strs:
                        rows.append(" | ".join(row_strs))
                pages.append({
                    "page": sheet_idx + 1,
                    "section": f"Sheet: {sheet_name}",
                    "content": "\n".join(rows)
                })
            return pages
        except Exception as e:
            return [{"page": 1, "section": "Error", "content": f"Excel extraction failed: {str(e)}"}]

    @staticmethod
    def _extract_csv(file_path: Path) -> List[Dict[str, Any]]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                lines = [" | ".join(row) for row in reader if any(row)]
            return [{
                "page": 1,
                "section": file_path.stem,
                "content": "\n".join(lines)
            }]
        except Exception as e:
            return [{"page": 1, "section": "Error", "content": f"CSV extraction failed: {str(e)}"}]

    @staticmethod
    def _extract_text(file_path: Path) -> List[Dict[str, Any]]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
            return [{
                "page": 1,
                "section": file_path.stem,
                "content": text
            }]
        except Exception as e:
            return [{"page": 1, "section": "Error", "content": f"Text extraction failed: {str(e)}"}]

    @staticmethod
    def _extract_email(file_path: Path) -> List[Dict[str, Any]]:
        """Extracts headers, body, and operational context from .eml files."""
        try:
            with open(file_path, "rb") as f:
                msg = email.message_from_binary_file(f, policy=policy.default)

            subject = str(msg.get("Subject", "No Subject"))
            sender = str(msg.get("From", "Unknown Sender"))
            date_str = str(msg.get("Date", ""))
            to_str = str(msg.get("To", ""))

            # Extract body (plain text or html stripped)
            body_text = ""
            body_part = msg.get_body(preferencelist=('plain', 'html'))
            if body_part:
                body_text = body_part.get_content()
                if body_part.get_content_type() == 'text/html':
                    body_text = re.sub(r'<[^>]+>', ' ', body_text)
                    body_text = re.sub(r'\s+', ' ', body_text).strip()

            header_summary = f"Subject: {subject}\nFrom: {sender}\nTo: {to_str}\nDate: {date_str}\n\n"
            full_content = header_summary + body_text

            return [{
                "page": 1,
                "section": f"Email: {subject[:60]}",
                "content": full_content,
                "extraction_mode": "email_parser",
                "email_metadata": {
                    "subject": subject,
                    "sender": sender,
                    "date": date_str,
                    "to": to_str
                }
            }]
        except Exception as e:
            return [{
                "page": 1,
                "section": "Error",
                "content": f"Email extraction failed: {str(e)}",
                "extraction_mode": "failed"
            }]

    @staticmethod
    def _extract_archive(file_path: Path) -> List[Dict[str, Any]]:
        """
        Extracts documents from a .zip archive with strict Zip Slip path traversal protection,
        size quotas, file count limits, and executable file rejection (never executed).
        """
        if not zipfile.is_zipfile(file_path):
            return [{
                "page": 1,
                "section": "Error",
                "content": f"File {file_path.name} is not a valid ZIP archive.",
                "extraction_mode": "failed"
            }]

        pages = []
        temp_extract_dir = tempfile.mkdtemp(prefix="intelgraph_zip_")
        try:
            with zipfile.ZipFile(file_path, "r") as zf:
                infolist = zf.infolist()

                # 1. Quota check: file count limit
                if len(infolist) > MAX_ARCHIVE_FILES:
                    return [{
                        "page": 1,
                        "section": "Quota Exceeded",
                        "content": f"Archive contains {len(infolist)} files (maximum limit is {MAX_ARCHIVE_FILES}). Extraction aborted for safety.",
                        "extraction_mode": "failed"
                    }]

                # 2. Quota check: uncompressed bytes limit
                total_uncompressed = sum(info.file_size for info in infolist)
                if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    return [{
                        "page": 1,
                        "section": "Quota Exceeded",
                        "content": f"Archive uncompressed size ({total_uncompressed / 1024 / 1024:.1f}MB) exceeds 50MB safety limit.",
                        "extraction_mode": "failed"
                    }]

                # 3. Zip Slip & Path Traversal check
                extract_dir_abs = os.path.abspath(temp_extract_dir)
                valid_members = []
                for member in infolist:
                    if member.is_dir():
                        continue

                    # Validate absolute target path does not escape sandbox
                    target_candidate = os.path.abspath(os.path.join(extract_dir_abs, member.filename))
                    if not target_candidate.startswith(extract_dir_abs + os.sep):
                        return [{
                            "page": 1,
                            "section": "Security Exception",
                            "content": f"Path traversal / Zip Slip attempt detected in archive member '{member.filename}'. Extraction aborted.",
                            "extraction_mode": "failed"
                        }]

                    # Check extension against executable rejection and allowed list
                    ext = Path(member.filename).suffix.lower()
                    if ext in DISALLOWED_EXECUTABLE_EXTENSIONS:
                        continue
                    if ext in ALLOWED_ARCHIVE_EXTENSIONS:
                        valid_members.append(member)

                # 4. Extract safe members only (never execute!)
                page_counter = 1
                for member in valid_members:
                    extracted_file_path = zf.extract(member, path=temp_extract_dir)
                    sub_ext = Path(extracted_file_path).suffix.lower().replace(".", "")
                    sub_pages = DocumentExtractor.extract(Path(extracted_file_path), sub_ext)
                    for sp in sub_pages:
                        pages.append({
                            "page": page_counter,
                            "section": f"{member.filename}: {sp.get('section', 'General')}",
                            "content": sp.get("content", ""),
                            "extraction_mode": "archive_member_extraction",
                            "archive_member": member.filename,
                            "ocr_confidence_pct": sp.get("ocr_confidence_pct")
                        })
                        page_counter += 1

            if not pages:
                return [{
                    "page": 1,
                    "section": "Notice",
                    "content": "Archive contained no supported document files (.pdf, .docx, .xlsx, .csv, .txt, .eml, .png).",
                    "extraction_mode": "archive_empty"
                }]
            return pages
        except Exception as e:
            return [{
                "page": 1,
                "section": "Error",
                "content": f"Archive processing failed: {str(e)}",
                "extraction_mode": "failed"
            }]
        finally:
            shutil.rmtree(temp_extract_dir, ignore_errors=True)

